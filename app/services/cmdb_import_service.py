"""Safe, snapshot-style import for the current CMDB/inventory domain."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Iterable

from sqlalchemy import insert, select

from app.database import SessionLocal, analyze_tables, engine
from app.models.cmdb_ci import CmdbCi
from app.models.cmdb_import_run import CmdbImportRun
from app.models.cmdb_staging_ci import CmdbStagingCi
from app.models.reconciliation_result import ReconciliationResult

_CHUNK_SIZE = 1_000

_ALIASES = {
    "hostname": "hostname", "machinename": "hostname", "machine_name": "hostname", "machine_name_": "hostname",
    "name": "hostname", "computername": "hostname", "computer_name": "hostname",
    "child": "hostname", "ven_installed": "ven_installed", "ven_installed_": "ven_installed",
    "pce_app": "cmdb_app", "app": "cmdb_app", "application": "cmdb_app",
    "pce_role": "cmdb_role", "role": "cmdb_role",
    "pce_env": "cmdb_env", "env": "cmdb_env", "environment": "cmdb_env",
    "pce_loc": "cmdb_loc", "loc": "cmdb_loc", "location": "cmdb_loc",
    "region": "region", "district": "district", "parent": "parent_application",
    "parent_sys_class_name": "parent_sys_class_name", "child_sys_class_name": "child_sys_class_name",
    "sys_created_on": "sys_created_on", "u_last_seen": "u_last_seen", "os": "os",
    "os_detail": "os_detail", "public_ip": "public_ip", "distinguished_name": "distinguished_name",
    "ip_with_default_gw": "ip_with_default_gw", "default_gw": "default_gw",
    "data_center": "data_center", "data_center_zone": "data_center_zone",
    "service_provider": "service_provider", "cloud_instance_id": "cloud_instance_id",
    "wave": "wave", "description": "description",
    # Stronger source identifiers are retained in extra_fields and used only
    # for duplicate accounting; active reconciliation still uses hostname.
    "sys_id": "source_identifier", "ci_sys_id": "source_identifier",
    "asset_id": "source_identifier", "source_id": "source_identifier",
}

_KNOWN_FIELDS = {
    "hostname", "ven_installed", "cmdb_app", "cmdb_role", "cmdb_env", "cmdb_loc",
    "region", "district", "wave", "description", "public_ip", "distinguished_name",
    "ip_with_default_gw", "default_gw", "data_center", "data_center_zone",
    "service_provider", "cloud_instance_id", "os", "os_detail", "parent_application",
    "parent_sys_class_name", "child_sys_class_name", "sys_created_on", "u_last_seen",
}

_ACTIVE_COLUMNS = tuple(column.name for column in CmdbCi.__table__.columns if column.name != "id")


def _utcnow():
    return datetime.now(UTC).replace(tzinfo=None)


def _header_key(value: str) -> str:
    return re.sub(r"[^a-z0-9.]+", "_", value.strip().casefold()).strip("_")


def _value(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_bool(value):
    normalized = (_value(value) or "").casefold()
    if normalized in {"true", "yes", "1", "t", "y"}:
        return True
    if normalized in {"false", "no", "0", "f", "n"}:
        return False
    return None


def _parse_dt(value):
    text = _value(value)
    if not text or text in {"None", "NaT"}:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


@dataclass(frozen=True)
class CmdbImportPreview:
    filename: str
    source_type: str
    headers: tuple[str, ...]
    recognized_headers: tuple[str, ...]
    unknown_headers: tuple[str, ...]
    identifier_strategy: str | None
    readable_rows: int
    malformed_rows: int
    duplicate_identities: int
    warnings: tuple[str, ...]
    errors: tuple[str, ...]

    @property
    def valid(self):
        return not self.errors


class CmdbValidationError(ValueError):
    def __init__(self, preview):
        self.preview = preview
        super().__init__("; ".join(preview.errors) or "CMDB validation failed")


class CmdbReplacementBlocked(RuntimeError):
    pass


class CmdbImportService:
    """Validate → stage → atomically activate a complete CMDB snapshot."""

    @classmethod
    def preview(cls, file_path) -> CmdbImportPreview:
        path = Path(file_path)
        if not path.is_file():
            return cls._empty_preview(path, "Selected file does not exist or is not a file.")
        try:
            headers = cls._headers(path)
            mapping, recognized, unknown = cls._header_mapping(headers)
            identity_strategy = "source identifier" if "source_identifier" in mapping else (
                "normalized hostname" if "hostname" in mapping else None
            )
            errors, warnings = [], []
            if identity_strategy is None:
                errors.append("No usable CMDB identifier column was detected (hostname/MachineName or source ID).")
            readable = malformed = duplicates = 0
            identities = set()
            for row in cls._iter_rows(path, headers):
                if not any(_value(value) for value in row.values()):
                    continue
                readable += 1
                identity = cls._identity(row, mapping)
                if identity is None:
                    malformed += 1
                    continue
                if identity in identities:
                    duplicates += 1
                identities.add(identity)
            if readable == 0:
                errors.append("File contains no readable CMDB rows.")
            if malformed:
                errors.append(f"{malformed} row(s) have no usable CMDB identity.")
            if duplicates:
                warnings.append(f"{duplicates} duplicate identity row(s) will be retained for later ambiguity review.")
            return CmdbImportPreview(
                path.name, cls._source_type(path), tuple(headers), tuple(sorted(recognized)), tuple(unknown),
                identity_strategy, readable, malformed, duplicates, tuple(warnings), tuple(errors),
            )
        except Exception as exc:
            return cls._empty_preview(path, f"Unable to parse {path.suffix.lower() or 'selected'} file: {exc}")

    @classmethod
    def import_file(cls, file_path, progress_callback=None):
        path = Path(file_path)
        preview = cls.preview(path)
        if not preview.valid:
            raise CmdbValidationError(preview)
        mapping, _, _ = cls._header_mapping(preview.headers)
        rows = (cls._row_to_mapping(row, mapping, _utcnow()) for row in cls._iter_rows(path, preview.headers))
        return cls._stage_and_activate(
            rows, preview, progress_callback=progress_callback,
        )

    @classmethod
    def replace_mock_records(cls, records: Iterable[dict], progress_callback=None):
        records = list(records)
        preview = CmdbImportPreview(
            filename="", source_type="mock_cmdb", headers=(), recognized_headers=(), unknown_headers=(),
            identifier_strategy="normalized hostname", readable_rows=len(records), malformed_rows=0,
            duplicate_identities=0, warnings=(), errors=(),
        )
        rows = (cls._mapping_from_record(record, _utcnow()) for record in records)
        return cls._stage_and_activate(rows, preview, progress_callback=progress_callback)

    @classmethod
    def _stage_and_activate(cls, rows, preview, progress_callback=None):
        cls._ensure_tables()
        session = SessionLocal()
        run = CmdbImportRun(
            source_type=preview.source_type, source_filename=preview.filename or None,
            status="STAGING", records_received=preview.readable_rows,
            duplicate_identities=preview.duplicate_identities,
            warning_message="\n".join(preview.warnings) or None,
        )
        session.add(run)
        session.commit()
        run_id = run.id
        session.close()

        staged = 0
        try:
            session = SessionLocal()
            buffer = []
            for row in rows:
                if not row:
                    continue
                row["import_run_id"] = run_id
                buffer.append(row)
                if len(buffer) >= _CHUNK_SIZE:
                    session.bulk_insert_mappings(CmdbStagingCi, buffer)
                    session.commit()
                    staged += len(buffer)
                    if progress_callback:
                        progress_callback(staged, preview.readable_rows)
                    buffer.clear()
            if buffer:
                session.bulk_insert_mappings(CmdbStagingCi, buffer)
                session.commit()
                staged += len(buffer)
            session.close()
        except Exception as exc:
            try:
                session.rollback()
                session.close()
            except Exception:
                pass
            cls._finish_failed_run(run_id, f"Staging failed: {exc}")
            cls._clear_staging(run_id)
            raise

        # A changed file between preview and staging is rejected before active
        # data changes. Validation's count is intentionally strict.
        if staged != preview.readable_rows:
            cls._finish_failed_run(run_id, "Staging row count did not match validated row count.")
            cls._clear_staging(run_id)
            raise CmdbValidationError(preview)

        session = SessionLocal()
        try:
            # Historical result rows contain immutable hostname/expected/current
            # values. Detach legacy navigation references before replacing the
            # active snapshot so those rows survive without blocking deletion.
            session.query(ReconciliationResult).filter(
                ReconciliationResult.cmdb_ci_id.isnot(None)
            ).update({"cmdb_ci_id": None}, synchronize_session=False)

            session.query(CmdbCi).delete(synchronize_session=False)
            staging_columns = [getattr(CmdbStagingCi, name) for name in _ACTIVE_COLUMNS]
            session.execute(insert(CmdbCi).from_select(
                _ACTIVE_COLUMNS,
                select(*staging_columns).where(CmdbStagingCi.import_run_id == run_id),
            ))
            session.query(CmdbStagingCi).filter(CmdbStagingCi.import_run_id == run_id).delete(
                synchronize_session=False
            )
            run = session.get(CmdbImportRun, run_id)
            run.status = "SUCCESS"
            run.records_inserted = staged
            run.completed_at = _utcnow()
            session.commit()
        except CmdbReplacementBlocked as exc:
            session.rollback()
            session.close()
            cls._finish_failed_run(run_id, str(exc), status="BLOCKED")
            cls._clear_staging(run_id)
            raise
        except Exception as exc:
            session.rollback()
            session.close()
            cls._finish_failed_run(run_id, f"Activation failed: {exc}")
            cls._clear_staging(run_id)
            raise
        else:
            session.close()

        analyze_tables("cmdb_cis")
        return {
            "status": "SUCCESS", "source_type": preview.source_type, "filename": preview.filename,
            "received": preview.readable_rows, "inserted": staged, "rejected": preview.malformed_rows,
            "duplicates": preview.duplicate_identities,
        }

    @staticmethod
    def _mapping_from_record(record, imported_at):
        mapping = {name: record.get(name) for name in _ACTIVE_COLUMNS}
        mapping["imported_at"] = imported_at
        mapping["extra_fields"] = record.get("extra_fields")
        return mapping

    @classmethod
    def _row_to_mapping(cls, row, mapping, imported_at):
        known, extra = {}, {}
        for header, value in row.items():
            canonical = _ALIASES.get(_header_key(header))
            value = _value(value)
            if canonical in _KNOWN_FIELDS:
                known[canonical] = value
            elif value is not None:
                extra[header] = value
        return {
            "hostname": known.get("hostname"), "cmdb_app": known.get("cmdb_app"),
            "cmdb_role": known.get("cmdb_role"), "cmdb_env": known.get("cmdb_env"),
            "cmdb_loc": known.get("cmdb_loc"), "ven_installed": _parse_bool(known.get("ven_installed")),
            "region": known.get("region"), "district": known.get("district"), "wave": known.get("wave"),
            "description": known.get("description"), "public_ip": known.get("public_ip"),
            "distinguished_name": known.get("distinguished_name"),
            "ip_with_default_gw": known.get("ip_with_default_gw"), "default_gw": known.get("default_gw"),
            "data_center": known.get("data_center"), "data_center_zone": known.get("data_center_zone"),
            "service_provider": known.get("service_provider"), "cloud_instance_id": known.get("cloud_instance_id"),
            "os": known.get("os"), "os_detail": known.get("os_detail"),
            "parent_application": known.get("parent_application"),
            "parent_sys_class_name": known.get("parent_sys_class_name"),
            "child_sys_class_name": known.get("child_sys_class_name"),
            "sys_created_on": _parse_dt(known.get("sys_created_on")),
            "u_last_seen": _parse_dt(known.get("u_last_seen")), "imported_at": imported_at,
            "extra_fields": extra or None,
        }

    @staticmethod
    def _identity(row, mapping):
        source_id = _value(row.get(mapping.get("source_identifier", "")))
        if source_id:
            return f"source:{source_id.casefold()}"
        hostname = _value(row.get(mapping.get("hostname", "")))
        return f"hostname:{hostname.casefold().rstrip('.')}" if hostname else None

    @staticmethod
    def _header_mapping(headers):
        mapping, known, unknown = {}, set(), []
        for header in headers:
            canonical = _ALIASES.get(_header_key(header))
            if canonical is None:
                unknown.append(header)
            else:
                known.add(canonical)
                mapping.setdefault(canonical, header)
        return mapping, known, unknown

    @staticmethod
    def _source_type(path):
        return {".csv": "inventory_csv", ".xlsx": "inventory_xlsx", ".xls": "inventory_xls"}.get(path.suffix.lower(), "unknown")

    @staticmethod
    def _headers(path):
        suffix = path.suffix.lower()
        if suffix == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                return tuple(header.strip() for header in (csv.reader(handle, strict=True).__next__()) if header)
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                return tuple(str(value).strip() for value in next(workbook.active.iter_rows(values_only=True)) if value is not None)
            finally:
                workbook.close()
        if suffix == ".xls":
            import importlib
            pd = importlib.import_module("pandas")
            return tuple(str(header).strip() for header in pd.read_excel(path, nrows=0).columns)
        raise ValueError(f"Unsupported file type: {suffix or 'none'}")

    @staticmethod
    def _iter_rows(path, headers):
        suffix = path.suffix.lower()
        if suffix == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                yield from csv.DictReader(handle, strict=True)
            return
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = workbook.active.iter_rows(values_only=True)
                next(rows)  # header row
                for values in rows:
                    yield dict(zip(headers, values))
            finally:
                workbook.close()
            return
        if suffix == ".xls":
            # pandas/xlrd currently materializes legacy XLS files. CSV and XLSX
            # are the memory-conscious paths; this remains supported with that limitation.
            import importlib
            pd = importlib.import_module("pandas")
            for row in pd.read_excel(path, dtype=str).fillna("").to_dict(orient="records"):
                yield row

    @staticmethod
    def _ensure_tables():
        CmdbImportRun.__table__.create(bind=engine, checkfirst=True)
        CmdbStagingCi.__table__.create(bind=engine, checkfirst=True)

    @staticmethod
    def _finish_failed_run(run_id, message, status="FAILED"):
        session = SessionLocal()
        try:
            run = session.get(CmdbImportRun, run_id)
            run.status = status
            run.error_message = message
            run.completed_at = _utcnow()
            session.commit()
        finally:
            session.close()

    @staticmethod
    def _clear_staging(run_id):
        session = SessionLocal()
        try:
            session.query(CmdbStagingCi).filter(CmdbStagingCi.import_run_id == run_id).delete(
                synchronize_session=False
            )
            session.commit()
        finally:
            session.close()

    @staticmethod
    def _empty_preview(path, error):
        return CmdbImportPreview(path.name, "unknown", (), (), (), None, 0, 0, 0, (), (error,))
